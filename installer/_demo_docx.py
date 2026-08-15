# coding=utf-8
"""数据集真实 docx 简历全流程演示：提取(段落+表格) → 清洗 → LLM 切片 → 向量化 → 检索"""
import json
import math
import os
import sys
import urllib.request

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "apps"))

from hr.services.resume_splitter import sanitize_resume_text, split_resume_text  # noqa: E402
from openai import OpenAI  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCX = os.path.join(ROOT, "testdata", "dataset_sample", "resume_sample_20200120", "docx", "cbb7a43eb62f.docx")
SENSE_KEY = os.environ.get("SENSENOVA_API_KEY", "")
SILICON_KEY = os.environ.get("SILICONFLOW_API_KEY", "")


def extract_docx(path):
    """当前生产实现（document.paragraphs）"""
    from docx import Document

    doc = Document(path)
    return chr(10).join(p.text for p in doc.paragraphs if p.text.strip())


def extract_docx_with_tables(path):
    """增强：段落 + 表格单元格"""
    from docx import Document

    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        seen = set()
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip() and cell._tc not in seen:
                    seen.add(cell._tc)
                    parts.append(cell.text.strip())
    return chr(10).join(parts)


print("=" * 70)
print("节点 0：数据集 docx 简历（cbb7a43eb62f.docx）")
print("-" * 70)
print("文件大小:", os.path.getsize(DOCX), "bytes")
text_prod = extract_docx(DOCX)
text_enhanced = extract_docx_with_tables(DOCX)
print("生产提取(document.paragraphs):", len(text_prod), "字符", "← 空！表格内容丢失" if not text_prod else "")
print("增强提取(段落+表格):", len(text_enhanced), "字符")

print()
print("=" * 70)
print("节点 1：提取文本（增强版，前 600 字符）")
print("-" * 70)
print(text_enhanced[:600].replace("\n", "⏎"))

# ---------- 节点 2：清洗 ----------
cleaned = sanitize_resume_text(text_enhanced)
print()
print("=" * 70)
print("节点 2：sanitize 清洗后 | 长度:", len(cleaned), "| 行数:", cleaned.count("\n") + 1)

# ---------- 节点 3：LLM 边界标注切片 ----------
llm = OpenAI(api_key=SENSE_KEY, base_url="https://token.sensenova.cn/v1")

def chat_fn(prompt):
    resp = llm.chat.completions.create(
        model="sensenova-6.8-flash-lite",
        messages=[{"role": "user", "content": prompt}],
        temperature=0.1, max_tokens=2000,
    )
    return resp.choices[0].message.content

print()
print("=" * 70)
print("节点 3：ResumeSplitter（LLM 行号边界标注，真实主路径）")
print("-" * 70)
print("输入行数:", len(cleaned.split("\n")), "| 前 3 行编号文本：")
for i, line in enumerate(cleaned.split("\n")[:3], 1):
    print("  %d  %s" % (i, line[:70]))
chunks = split_resume_text(cleaned, chat_fn)
print("→ 切片结果:", len(chunks), "段")
for i, chunk in enumerate(chunks):
    pii = "[已脱敏]" in chunk["content"]
    print("  [%d] title=%r (%d字)%s %s..." % (
        i + 1, chunk["title"], len(chunk["content"]), " PII掩码✓" if pii else "",
        chunk["content"][:60].replace("\n", "⏎")))

# ---------- 节点 4：向量化 ----------
def embed_one(text):
    req = urllib.request.Request(
        "https://api.siliconflow.cn/v1/embeddings",
        data=json.dumps({"model": "BAAI/bge-large-zh-v1.5", "input": [text]}).encode(),
        headers={"Content-Type": "application/json", "Authorization": "Bearer " + SILICON_KEY},
    )
    with urllib.request.urlopen(req, timeout=60) as resp:
        return json.loads(resp.read().decode())["data"][0]["embedding"]

print()
print("=" * 70)
print("节点 4：向量化（bge-large-zh-v1.5，1024 维）")
print("-" * 70)
vecs = []
for i, chunk in enumerate(chunks):
    vec = embed_one(chunk["content"])
    vecs.append(vec)
    norm = math.sqrt(sum(x * x for x in vec))
    print("  段[%d] 维度=%d 前5维=[%s] 范数=%.4f" % (
        i + 1, len(vec), ", ".join("%.4f" % x for x in vec[:5]), norm))

# ---------- 节点 5：检索 ----------
query = "有 Unity 场景搭建和灯光渲染经验的候选人"
print()
print("=" * 70)
print("节点 5：检索（查询 → 余弦排序）")
print("-" * 70)
print("查询:", query)
qv = embed_one(query)

def cos(a, b):
    return sum(x * y for x, y in zip(a, b)) / (math.sqrt(sum(x * x for x in a)) * math.sqrt(sum(y * y for y in b)))

ranked = sorted(range(len(chunks)), key=lambda i: -cos(qv, vecs[i]))
for rank, i in enumerate(ranked[:3], 1):
    print("  #%d 相似度=%.4f %s..." % (rank, cos(qv, vecs[i]), chunks[i]["content"][:60].replace("\n", "⏎")))
print()
print("全流程：提取 %d 字 → 清洗 %d 字 → 切片 %d 段 → 向量 %d 条 → 检索" % (
    len(text_enhanced), len(cleaned), len(chunks), len(vecs)))
