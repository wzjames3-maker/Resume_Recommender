import json, os

API = r"C:\Users\Administrator\Desktop\简历推荐系统\src\frontend\api_client.py"

# Check if already done
with open(API, "r", encoding="utf-8") as f:
    existing = f.read()
if "def get_resume(" in existing:
    print("SKIP: already appended")
else:
    # New code to append
    newcode = '''
# ============================================================
# v2.0 NEW: Resume CRUD + Stats + Export (T-036)  
# ============================================================
import os as _os
from pymongo import MongoClient
from io import BytesIO

MONGO_URI = _os.getenv("MONGODB_URL", "mongodb://admin:password@mongodb:27017/resume_rag?authSource=admin")

@st.cache_resource
def _get_mongo_client():
    return MongoClient(MONGO_URI)

def _get_resume_collection():
    return _get_mongo_client()["resume_rag"]["resumes"]

def get_resume(resume_id):
    try:
        col = _get_resume_collection()
        doc = col.find_one({"id": resume_id})
        if doc:
            doc.pop("_id", None)
            return doc
        return None
    except Exception as e:
        st.error(f"获取简历失败: {e}")
        return None

def list_resumes_api(keyword="", skill="", city="", education="", page=1, size=20):
    try:
        col = _get_resume_collection()
        query = {"status": {"$ne": "deleted"}}
        if keyword:
            query["$or"] = [
                {"personal_info.name": {"$regex": keyword, "$options": "i"}},
                {"skill_list": {"$regex": keyword, "$options": "i"}},
            ]
        if skill:
            query["skill_list"] = {"$regex": skill, "$options": "i"}
        if city:
            query["personal_info.expected_cities"] = {"$regex": city, "$options": "i"}
        if education:
            query["personal_info.highest_education"] = {"$regex": education, "$options": "i"}
        total = col.count_documents(query)
        docs = list(col.find(query).sort("created_at", -1).skip((page - 1) * size).limit(size))
        for d in docs:
            d.pop("_id", None)
        return {"items": docs, "total": total, "page": page, "size": size}
    except Exception as e:
        st.error(f"获取简历列表失败: {e}")
        return None

def update_resume(resume_id, data):
    try:
        col = _get_resume_collection()
        result = col.update_one({"id": resume_id}, {"$set": data})
        return result.modified_count > 0
    except Exception as e:
        st.error(f"更新失败: {e}")
        return False

def delete_resume_api(resume_id):
    try:
        col = _get_resume_collection()
        result = col.update_one({"id": resume_id}, {"$set": {"status": "deleted"}})
        return result.modified_count > 0
    except Exception as e:
        st.error(f"删除失败: {e}")
        return False

def get_resume_stats():
    try:
        col = _get_resume_collection()
        skills = list(col.aggregate([
            {"$match": {"status": {"$ne": "deleted"}}},
            {"$unwind": "$skill_list"},
            {"$group": {"_id": "$skill_list", "count": {"$sum": 1}}},
            {"$sort": {"count": -1}}, {"$limit": 50}
        ]))
        educations = list(col.aggregate([
            {"$match": {"status": {"$ne": "deleted"}, "personal_info.highest_education": {"$exists": True}}},
            {"$group": {"_id": "$personal_info.highest_education", "count": {"$sum": 1}}},
            {"$sort": {"count": -1}}
        ]))
        cities = list(col.aggregate([
            {"$match": {"status": {"$ne": "deleted"}}},
            {"$unwind": "$personal_info.expected_cities"},
            {"$group": {"_id": "$personal_info.expected_cities", "count": {"$sum": 1}}},
            {"$sort": {"count": -1}}, {"$limit": 20}
        ]))
        companies = list(col.aggregate([
            {"$match": {"status": {"$ne": "deleted"}}},
            {"$unwind": "$experience_list"},
            {"$group": {"_id": "$experience_list.company", "count": {"$sum": 1}}},
            {"$sort": {"count": -1}}, {"$limit": 20}
        ]))
        total = col.count_documents({"status": {"$ne": "deleted"}})
        return {"total_resumes": total, "skills": skills, "educations": educations, "cities": cities, "companies": companies}
    except Exception as e:
        st.error(f"获取统计失败: {e}")
        return None

def export_candidates_excel(candidates, filename="candidates.xlsx"):
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "results"
        hf = Font(bold=True, color="FFFFFF", size=11)
        hfill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        ha = Alignment(horizontal="center", vertical="center")
        tb = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
        hdrs = ["Rank", "Name", "Score", "Matched Skills", "Missing Skills"]
        for col, h in enumerate(hdrs, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = hf; cell.fill = hfill; cell.alignment = ha; cell.border = tb
        for i, c in enumerate(candidates):
            row = i + 2
            meta = c.get("metadata", {})
            reason = c.get("reason", {})
            ws.cell(row=row, column=1, value=i+1).border = tb
            ws.cell(row=row, column=2, value=meta.get("name", "-")).border = tb
            ws.cell(row=row, column=3, value=f"{c.get('score', 0):.1%}").border = tb
            ws.cell(row=row, column=4, value=", ".join(reason.get("matched_skills", []))).border = tb
            ws.cell(row=row, column=5, value=", ".join(reason.get("missing_skills", []))).border = tb
        for col, w in enumerate([6, 12, 10, 30, 30], 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w
        out = BytesIO(); wb.save(out); out.seek(0)
        return out.getvalue()
    except Exception as e:
        st.error(f"导出失败: {e}")
        return None
'''
    with open(API, "a", encoding="utf-8") as f:
        f.write(newcode)
    print("OK: api_client.py appended with 6 new methods")
