# ====== Append new methods to api_client.py ======
api_client_path = r"C:\Users\Administrator\Desktop\简历推荐系统\src\frontend\api_client.py"
with open(api_client_path, "r", encoding="utf-8") as f:
    current = f.read()

new_methods = r"""

# ============================================================
# v2.0 NEW: Resume management & Export methods (T-036)
# ============================================================

import os
from pymongo import MongoClient
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from io import BytesIO

MONGO_URI = os.getenv("MONGODB_URL", "mongodb://admin:password@mongodb:27017/resume_rag?authSource=admin")

@st.cache_resource
def _get_mongo_client():
    return MongoClient(MONGO_URI)

def _get_resume_collection():
    return _get_mongo_client()["resume_rag"]["resumes"]


def get_resume(resume_id):
    """Get resume detail from MongoDB."""
    try:
        col = _get_resume_collection()
        doc = col.find_one({"id": resume_id})
        if doc:
            doc.pop("_id", None)
            return doc
        return None
    except Exception as e:
        st.error(f"获取简历详情失败: {e}")
        return None


def list_resumes_api(
    keyword="",
    skill="",
    city="",
    education="",
    min_experience="",
    is_985=False,
    is_211=False,
    status=None,
    page=1,
    size=20,
):
    """List resumes with advanced filters from MongoDB."""
    try:
        col = _get_resume_collection()
        query = {"status": {"$ne": "deleted"}}
        if keyword:
            query["$or"] = [
                {"personal_info.name": {"$regex": keyword, "$options": "i"}},
                {"skill_list": {"$regex": keyword, "$options": "i"}},
                {"raw_text": {"$regex": keyword, "$options": "i"}},
            ]
        if skill:
            query["skill_list"] = {"$regex": skill, "$options": "i"}
        if city:
            query["personal_info.expected_cities"] = {"$regex": city, "$options": "i"}
        if education:
            query["personal_info.highest_education"] = {"$regex": education, "$options": "i"}
        if min_experience:
            query["personal_info.total_experience"] = {"$regex": min_experience, "$options": "i"}
        if is_985:
            query["education_list.is_985"] = True
        if is_211:
            query["education_list.is_211"] = True
        if status:
            query["status"] = status

        total = col.count_documents(query)
        docs = list(col.find(query).sort("created_at", -1).skip((page - 1) * size).limit(size))
        for d in docs:
            d.pop("_id", None)
        return {"items": docs, "total": total, "page": page, "size": size}
    except Exception as e:
        st.error(f"获取简历列表失败: {e}")
        return None


def update_resume(resume_id, data):
    """Update resume fields."""
    try:
        col = _get_resume_collection()
        result = col.update_one({"id": resume_id}, {"$set": data})
        return result.modified_count > 0
    except Exception as e:
        st.error(f"更新简历失败: {e}")
        return False


def delete_resume_api(resume_id):
    """Soft-delete a resume."""
    try:
        col = _get_resume_collection()
        result = col.update_one({"id": resume_id}, {"$set": {"status": "deleted"}})
        return result.modified_count > 0
    except Exception as e:
        st.error(f"删除简历失败: {e}")
        return False


def get_resume_stats():
    """Get aggregated statistics from MongoDB."""
    try:
        col = _get_resume_collection()
        # Skill distribution
        skill_pipeline = [
            {"$match": {"status": {"$ne": "deleted"}}},
            {"$unwind": "$skill_list"},
            {"$group": {"_id": "$skill_list", "count": {"$sum": 1}}},
            {"$sort": {"count": -1}},
            {"$limit": 50},
        ]
        skills = list(col.aggregate(skill_pipeline))

        # Education distribution
        edu_pipeline = [
            {"$match": {"status": {"$ne": "deleted"}, "personal_info.highest_education": {"$exists": True}}},
            {"$group": {"_id": "$personal_info.highest_education", "count": {"$sum": 1}}},
            {"$sort": {"count": -1}},
        ]
        educations = list(col.aggregate(edu_pipeline))

        # City distribution
        city_pipeline = [
            {"$match": {"status": {"$ne": "deleted"}}},
            {"$unwind": "$personal_info.expected_cities"},
            {"$group": {"_id": "$personal_info.expected_cities", "count": {"$sum": 1}}},
            {"$sort": {"count": -1}},
            {"$limit": 20},
        ]
        cities = list(col.aggregate(city_pipeline))

        # Company distribution
        company_pipeline = [
            {"$match": {"status": {"$ne": "deleted"}}},
            {"$unwind": "$experience_list"},
            {"$group": {"_id": "$experience_list.company", "count": {"$sum": 1}}},
            {"$sort": {"count": -1}},
            {"$limit": 20},
        ]
        companies = list(col.aggregate(company_pipeline))

        # Total count
        total = col.count_documents({"status": {"$ne": "deleted"}})

        return {
            "total_resumes": total,
            "skills": skills,
            "educations": educations,
            "cities": cities,
            "companies": companies,
        }
    except Exception as e:
        st.error(f"获取统计数据失败: {e}")
        return None


def export_candidates_excel(candidates, filename="candidates.xlsx"):
    """Export candidates list to Excel bytes."""
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "候选人推荐结果"

        # Header style
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        headers = ["排名", "姓名", "匹配度", "匹配技能", "缺失技能", "学历", "毕业院校", "工作年限", "当前状态"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        for i, c in enumerate(candidates):
            row = i + 2
            meta = c.get("metadata", {})
            reason = c.get("reason", {})

            ws.cell(row=row, column=1, value=i + 1).border = thin_border
            ws.cell(row=row, column=2, value=meta.get("name", "-")).border = thin_border
            ws.cell(row=row, column=3, value=f"{c.get('score', 0):.1%}").border = thin_border
            ws.cell(row=row, column=4, value=", ".join(reason.get("matched_skills", []))).border = thin_border
            ws.cell(row=row, column=5, value=", ".join(reason.get("missing_skills", []))).border = thin_border

            # Get resume detail for extra info
            resume_id = c.get("resume_id")
            if resume_id:
                doc = get_resume(resume_id)
                if doc:
                    pi = doc.get("personal_info", {})
                    edu = doc.get("education_list", [{}])[0] if doc.get("education_list") else {}
                    ws.cell(row=row, column=6, value=pi.get("highest_education", "-")).border = thin_border
                    ws.cell(row=row, column=7, value=edu.get("school", "-")).border = thin_border
                    ws.cell(row=row, column=8, value=pi.get("total_experience", "-")).border = thin_border
                    ws.cell(row=row, column=9, value=pi.get("current_status", "-")).border = thin_border
                else:
                    for ccol in range(6, 10):
                        ws.cell(row=row, column=ccol, value="-").border = thin_border
            else:
                for ccol in range(6, 10):
                    ws.cell(row=row, column=ccol, value="-").border = thin_border

        # Adjust column widths
        widths = [6, 12, 10, 25, 25, 10, 15, 10, 10]
        for col, w in enumerate(widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()
    except Exception as e:
        st.error(f"导出Excel失败: {e}")
        return None
"""

with open(api_client_path, "a", encoding="utf-8") as f:
    f.write(new_methods)

print("api_client.py: Added 6 new methods")
