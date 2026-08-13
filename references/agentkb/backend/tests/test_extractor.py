
from app.services.resume.blocks import BlockKind
from app.services.resume.extractor import (
    DocxExtractor,
    PdfExtractor,
    TextExtractor,
    extract_resume,
)


def _build_docx_bytes() -> bytes:
    # 用 python-docx 构造：一个段落 + 一个表格（含「标签:值」）+ 一个文本框
    import io

    from docx import Document as Docx
    from docx.oxml import OxmlElement

    doc = Docx()
    doc.add_paragraph("张三个人简历")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "姓名"
    table.cell(0, 1).text = "张三"
    table.cell(1, 0).text = "电话"
    table.cell(1, 1).text = "13800000000"
    # 文本框：手工插入 w:txbxContent
    p = doc.add_paragraph()
    run = p.add_run()
    txbx = OxmlElement("w:txbxContent")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "文本框里的技能：Python"
    r.append(t)
    txbx.append(r)
    run._r.append(OxmlElement("w:drawing"))
    body = doc.element.body
    drawing = OxmlElement("w:drawing")
    drawing.insert(0, txbx)
    body.insert(0, drawing)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_extract_contains_paragraph_and_table():
    data = _build_docx_bytes()
    blocks = DocxExtractor().extract(data)
    text = "\n".join(b.content for b in blocks)
    assert "张三个人简历" in text          # 段落通道
    assert "姓名" in text and "电话" in text  # 表格通道行主序
    # 表格行展平保留「标签:值」：至少含「姓名」与「张三」
    assert "张三" in text
    # 文本框通道（w:txbxContent XML 遍历）
    textbox_blocks = [b for b in blocks if b.kind == BlockKind.textbox]
    assert textbox_blocks and "Python" in textbox_blocks[0].content


def test_docx_extract_preserves_label_value_structure():
    data = _build_docx_bytes()
    blocks = DocxExtractor().extract(data)
    table_blocks = [b for b in blocks if b.kind == BlockKind.table]
    assert table_blocks, "应产出表格块"
    row_block = table_blocks[0].content
    assert "姓名" in row_block and "张三" in row_block


def _build_pdf_bytes() -> bytes:
    # 用 pypdf 构造带 Helvetica 字体资源的文本页；手工裸 PDF 无 xref/%%EOF，pypdf
    # 在 6.x 下无法可靠重建流边界（原计划 `_MINIMAL_PDF` 写法在新版本必挂）。
    import io

    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=595, height=842)
    font = DictionaryObject()
    font.update({NameObject("/Type"): NameObject("/Font"),
                 NameObject("/Subtype"): NameObject("/Type1"),
                 NameObject("/BaseFont"): NameObject("/Helvetica")})
    font_ref = writer._add_object(font)
    resources = DictionaryObject()
    resources.update({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})})
    page[NameObject("/Resources")] = resources
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 72 720 Td (PDF resume test) Tj ET")
    stream_ref = writer._add_object(content)
    page[NameObject("/Contents")] = stream_ref
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def test_pdf_extract_text_layer(tmp_path):
    path = tmp_path / "resume.pdf"
    path.write_bytes(_build_pdf_bytes())
    blocks = PdfExtractor().extract(str(path))
    assert "PDF resume test" in "\n".join(b.content for b in blocks)


def test_text_extract():
    # str 解释为文件路径、bytes 解释为内容（与 extract_resume 契约一致）
    blocks = TextExtractor().extract("张三\n电话 13800000000".encode())
    assert blocks[0].content == "张三"


def test_extract_resume_dispatch_by_format(tmp_path):
    (tmp_path / "a.txt").write_text("张三", encoding="utf-8")
    (tmp_path / "a.docx").write_bytes(_build_docx_bytes())
    p = tmp_path / "a.pdf"
    import pypdf
    w = pypdf.PdfWriter()
    w.add_blank_page(200, 200)
    with open(p, "wb") as f:
        w.write(f)
    assert extract_resume(str(tmp_path / "a.txt"), "txt")[0].content == "张三"
    assert extract_resume(str(tmp_path / "a.docx"), "docx")[0].kind == BlockKind.paragraph
    assert extract_resume(str(p), "pdf") is not None