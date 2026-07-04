"""
智能招聘 RAG 推荐系统 - API 客户端封装

所有前端页面共用的 API 调用逻辑
"""

import requests
import json
import os
import streamlit as st
from typing import Optional, Dict, Any, List

API_BASE_URL = os.environ.get("API_BASE_URL", "http://localhost/api/v1")


def get_token() -> Optional[str]:
    return st.session_state.get("token")


def get_headers() -> Dict[str, str]:
    token = get_token()
    if token:
        return {"Authorization": f"Bearer {token}"}
    return {}


def login(username: str, password: str) -> bool:
    try:
        r = requests.post(f"{API_BASE_URL}/auth/login", json={"username": username, "password": password})
        if r.status_code == 200:
            data = r.json()
            st.session_state.token = data.get("token")
            st.session_state.user_id = data.get("user_id")
            st.session_state.role = data.get("role")
            return True
        return False
    except Exception as e:
        st.error(f"登录失败: {e}")
        return False




def verify_token(token: str) -> Optional[Dict]:
    """验证 JWT Token 是否有效"""
    try:
        r = requests.post(
            f"{API_BASE_URL}/auth/verify",
            json={"token": token},
            timeout=10,
        )
        if r.status_code == 200:
            data = r.json()
            if data.get("valid"):
                return data
        return None
    except Exception:
        return None


def logout():
    for key in ["token", "user_id", "role", "conversation_id", "messages"]:
        if key in st.session_state:
            del st.session_state[key]


def health_check() -> Optional[Dict]:
    try:
        health_url = f"{API_BASE_URL.replace('/api/v1', '')}/health"
        r = requests.get(health_url, timeout=5)
        return r.json() if r.status_code == 200 else None
    except Exception:
        return None


def upload_resume(file_bytes: bytes, filename: str) -> Optional[Dict]:
    try:
        r = requests.post(
            f"{API_BASE_URL}/resumes/upload",
            files={"file": (filename, file_bytes)},
            headers=get_headers(),
            timeout=120,
        )
        return r.json() if r.status_code == 200 else None
    except Exception as e:
        return {"success": False, "message": str(e)}


def send_chat_message(message: str, conversation_id: Optional[str] = None) -> Optional[Dict]:
    try:
        payload = {"message": message}
        if conversation_id:
            payload["conversation_id"] = conversation_id

        r = requests.post(
            f"{API_BASE_URL}/chat",
            json=payload,
            headers=get_headers(),
            stream=True,
            timeout=60,
        )

        if r.status_code != 200:
            return None

        full_response = ""
        candidates = []
        session_id = None

        for line in r.iter_lines():
            if line:
                line = line.decode("utf-8")
                if line.startswith("data: "):
                    data = json.loads(line[6:])
                    if data.get("type") == "token":
                        full_response += data.get("content", "")
                    elif data.get("type") == "sources":
                        candidates = data.get("candidates", [])
                    elif data.get("type") == "done":
                        session_id = data.get("session_id")

        return {"message": full_response, "candidates": candidates, "session_id": session_id}
    except Exception as e:
        st.error(f"发送消息失败: {e}")
        return None


def list_conversations(page: int = 1, size: int = 20) -> Optional[Dict]:
    try:
        r = requests.get(
            f"{API_BASE_URL}/conversations",
            params={"page": page, "size": size},
            headers=get_headers(),
            timeout=10,
        )
        return r.json() if r.status_code == 200 else None
    except Exception:
        return None


def get_conversation(conversation_id: str) -> Optional[Dict]:
    try:
        r = requests.get(
            f"{API_BASE_URL}/conversations/{conversation_id}",
            headers=get_headers(),
            timeout=10,
        )
        return r.json() if r.status_code == 200 else None
    except:
        return None


def delete_conversation(conversation_id: str) -> bool:
    try:
        r = requests.delete(
            f"{API_BASE_URL}/conversations/{conversation_id}",
            headers=get_headers(),
            timeout=10,
        )
        return r.status_code == 200
    except Exception:
        return False

# ============================================================
# Resume CRUD + Stats + Export (通过 API 调用，不再直连 MongoDB)
# ============================================================
from io import BytesIO


def get_resume(resume_id):
    """获取简历详情"""
    try:
        r = requests.get(
            f"{API_BASE_URL}/resumes/{resume_id}",
            headers=get_headers(),
            timeout=10,
        )
        return r.json() if r.status_code == 200 else None
    except Exception as e:
        st.error(f"获取简历失败: {e}")
        return None


def list_resumes_api(keyword="", skill="", city="", education="", min_exp=0, is_985=False, is_211=False, page=1, size=20):
    """获取简历列表"""
    try:
        r = requests.get(
            f"{API_BASE_URL}/resumes/",
            params={"keyword": keyword, "skill": skill, "city": city,
                    "education": education, "min_exp": min_exp,
                    "is_985": is_985, "is_211": is_211,
                    "page": page, "size": size},
            headers=get_headers(),
            timeout=10,
        )
        return r.json() if r.status_code == 200 else None
    except Exception as e:
        st.error(f"获取简历列表失败: {e}")
        return None


def update_resume(resume_id, data):
    """更新简历"""
    try:
        r = requests.put(
            f"{API_BASE_URL}/resumes/{resume_id}",
            json=data,
            headers=get_headers(),
            timeout=10,
        )
        if r.status_code == 200:
            return True
        st.error(f"更新失败: {r.json().get('detail', r.text)}")
        return False
    except Exception as e:
        st.error(f"更新失败: {e}")
        return False


def delete_resume_api(resume_id):
    """软删除简历"""
    try:
        r = requests.delete(
            f"{API_BASE_URL}/resumes/{resume_id}",
            headers=get_headers(),
            timeout=10,
        )
        if r.status_code == 200:
            return True
        st.error(f"删除失败: {r.json().get('detail', r.text)}")
        return False
    except Exception as e:
        st.error(f"删除失败: {e}")
        return False


def get_resume_stats():
    """获取简历统计信息"""
    try:
        r = requests.get(
            f"{API_BASE_URL}/resumes/stats",
            headers=get_headers(),
            timeout=10,
        )
        return r.json() if r.status_code == 200 else None
    except Exception as e:
        st.error(f"获取统计失败: {e}")
        return None


def export_candidates_excel(candidates, filename="candidates.xlsx"):
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "results"
        hf = Font(bold=True, color="FFFFFF", size=11)
        hfill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        ha = Alignment(horizontal="center", vertical="center")
        tb = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
        hdrs = ["Rank", "Name", "Score", "Matched Skills", "Missing Skills"]
        for col, h in enumerate(hdrs, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = hf; cell.fill = hfill; cell.alignment = ha; cell.border = tb
        for i, c in enumerate(candidates):
            row = i + 2
            meta = c.get("metadata", {})
            reason = c.get("reason", {})
            ws.cell(row=row, column=1, value=i+1).border = tb
            ws.cell(row=row, column=2, value=meta.get("name", "-")).border = tb
            ws.cell(row=row, column=3, value=f"{c.get('score', 0):.1%}").border = tb
            ws.cell(row=row, column=4, value=", ".join(reason.get("matched_skills", []))).border = tb
            ws.cell(row=row, column=5, value=", ".join(reason.get("missing_skills", []))).border = tb
        for col, w in enumerate([6, 12, 10, 30, 30], 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w
        out = BytesIO(); wb.save(out); out.seek(0)
        return out.getvalue()
    except Exception as e:
        st.error(f"导出失败: {e}")
        return None

