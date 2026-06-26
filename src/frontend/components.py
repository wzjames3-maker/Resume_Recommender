"""
智能招聘 RAG 推荐系统 - 候选人卡片组件

Streamlit 候选人卡片展示
"""

import streamlit as st
from typing import Any, Dict, List, Optional


def render_candidate_card(candidate: Dict[str, Any], index: int):
    """
    渲染候选人卡片

    Args:
        candidate: 候选人数据
        index: 索引
    """
    with st.container():
        # 卡片头部
        col1, col2 = st.columns([3, 1])

        with col1:
            st.markdown(f"### 候选人 #{index + 1}")
            st.markdown(f"**ID:** {candidate.get('resume_id', 'N/A')}")

        with col2:
            score = candidate.get('score', 0)
            st.metric("匹配度", f"{score:.1%}")

        # 推荐理由
        reason_data = candidate.get('reason', {})
        if reason_data:
            st.markdown("**推荐理由：**")
            st.info(reason_data.get('reason', '暂无'))

            # 技能匹配
            col1, col2 = st.columns(2)

            with col1:
                matched_skills = reason_data.get('matched_skills', [])
                if matched_skills:
                    st.markdown("**✅ 匹配技能：**")
                    for skill in matched_skills:
                        st.markdown(f"- {skill}")

            with col2:
                missing_skills = reason_data.get('missing_skills', [])
                if missing_skills:
                    st.markdown("**❌ 缺失技能：**")
                    for skill in missing_skills:
                        st.markdown(f"- {skill}")

        # 候选人内容摘要
        content = candidate.get('content', '')
        if content:
            with st.expander("查看详情"):
                st.markdown(content[:500] + "..." if len(content) > 500 else content)

        # Score Breakdown
        score_breakdown = reason_data.get('score_breakdown', {})
        if score_breakdown:
            with st.expander("分数详情"):
                for key, value in score_breakdown.items():
                    st.markdown(f"- **{key}:** {value:.3f}")

        st.markdown("---")


def render_candidate_list(candidates: List[Dict[str, Any]]):
    """
    渲染候选人列表

    Args:
        candidates: 候选人列表
    """
    if not candidates:
        st.info("暂无推荐候选人")
        return

    st.markdown(f"### 📋 推荐候选人（共 {len(candidates)} 位）")

    for i, candidate in enumerate(candidates):
        render_candidate_card(candidate, i)


def render_sidebar_stats(candidates: List[Dict[str, Any]]):
    """
    渲染侧边栏统计信息

    Args:
        candidates: 候选人列表
    """
    with st.sidebar:
        st.markdown("---")
        st.markdown("### 📊 统计信息")

        if candidates:
            # 计算平均匹配度
            scores = [c.get('score', 0) for c in candidates]
            avg_score = sum(scores) / len(scores) if scores else 0

            st.metric("候选人数量", len(candidates))
            st.metric("平均匹配度", f"{avg_score:.1%}")

            # 技能统计
            all_matched_skills = []
            all_missing_skills = []

            for candidate in candidates:
                reason = candidate.get('reason', {})
                all_matched_skills.extend(reason.get('matched_skills', []))
                all_missing_skills.extend(reason.get('missing_skills', []))

            if all_matched_skills:
                st.markdown("**常见匹配技能：**")
                skill_counts = {}
                for skill in all_matched_skills:
                    skill_counts[skill] = skill_counts.get(skill, 0) + 1

                for skill, count in sorted(skill_counts.items(), key=lambda x: x[1], reverse=True)[:5]:
                    st.markdown(f"- {skill} ({count})")

            if all_missing_skills:
                st.markdown("**常见缺失技能：**")
                skill_counts = {}
                for skill in all_missing_skills:
                    skill_counts[skill] = skill_counts.get(skill, 0) + 1

                for skill, count in sorted(skill_counts.items(), key=lambda x: x[1], reverse=True)[:5]:
                    st.markdown(f"- {skill} ({count})")
        else:
            st.info("暂无统计数据")


# ============================================================
# v2.0 NEW: Full resume display, filter sidebar, stats card, Excel export (T-036)
# ============================================================

from io import BytesIO


def render_resume_full(doc, show_id=True):
    """Render a complete resume with all sections."""
    if not doc:
        st.warning("简历数据为空")
        return
    pi = doc.get("personal_info", {})
    resume_id = doc.get("id", "")
    if show_id:
        st.caption(f"ID: `{resume_id}`")
    st.markdown(f"## {pi.get('full_name', '未知')}")
    col1, col2, col3, col4 = st.columns(4)
    with col1:
        st.metric("学历", pi.get("highest_education") or (doc.get("education_list", [{}])[0].get("degree", "") if doc.get("education_list") else "") or "-")
    with col2:
        st.metric("工作年限", f"{pi.get('years_of_experience', '-')} 年")
    with col3:
        cities = pi.get("expected_city")
        city_str = ", ".join(cities) if isinstance(cities, list) and cities else str(cities or "-")
        st.metric("期望城市", city_str)
    with col4:
        st.metric("期望薪资", (pi.get("expected_salary_range") or "-"))
    st.markdown("---")
    skills = doc.get("skill_list", [])
    if skills:
        st.markdown("### 技能标签")
        tags_html = " ".join([
            f'<span style="display:inline-block;background:#e0e7ff;color:#3730a3;padding:3px 10px;margin:3px;border-radius:12px;font-size:13px;">{s if isinstance(s, str) else s.get("name", str(s))}</span>'
            for s in skills[:30]
        ])
        st.markdown(tags_html, unsafe_allow_html=True)
        if len(skills) > 30:
            st.caption(f"...及其他 {len(skills) - 30} 项技能")
        st.markdown("---")
    tab1, tab2, tab3 = st.tabs(["教育经历", "工作经历", "项目经历"])
    with tab1:
        edu_list = doc.get("education_list", [])
        if edu_list:
            for edu in edu_list:
                school = edu.get("school", "-")
                degree = edu.get("degree", "")
                major = edu.get("major", "")
                is_985 = edu.get("is_985", False)
                is_211 = edu.get("is_211", False)
                badges = []
                if is_985: badges.append("985")
                if is_211: badges.append("211")
                badge_str = " ".join([f"`{b}`" for b in badges]) if badges else ""
                st.markdown(f"**{school}** {badge_str}")
                st.caption(f"{degree} | {major}")
                st.markdown("---")
        else:
            st.info("暂无教育经历")
    with tab2:
        exp_list = doc.get("experience_list", [])
        if exp_list:
            for exp in exp_list:
                company = exp.get("company", "-")
                title = exp.get("title", "-")
                duration = exp.get("start_date", "-")
                desc = exp.get("description", "")
                st.markdown(f"**{company}** — {title}")
                st.caption(f"📅 {duration}")
                if desc:
                    st.markdown(desc[:300] + ("..." if len(desc) > 300 else ""))
                st.markdown("---")
        else:
            st.info("暂无工作经历")
    with tab3:
        proj_list = doc.get("project_list", [])
        if proj_list:
            for proj in proj_list:
                name = proj.get("name", "-")
                role = proj.get("role", "-")
                tech_stack = proj.get("tech_stack", [])
                desc = proj.get("description", "")
                st.markdown(f"**{name}** — {role}")
                if tech_stack:
                    tags = " ".join([f"`{t}`" for t in tech_stack])
                    st.markdown(tags)
                if desc:
                    st.caption(desc[:300] + ("..." if len(desc) > 300 else ""))
                st.markdown("---")
        else:
            st.info("暂无项目经历")


def render_filter_sidebar(cities=None, skills=None):
    """Render advanced filter panel. Call inside a 'with st.sidebar:' context."""
    st.markdown("---")
    st.markdown("### 🔍 高级筛选")
    filters = {}
    if cities:
        filters["city"] = st.multiselect("期望城市", options=cities, default=[], help="选择期望工作城市")
    else:
        filters["city"] = st.text_input("期望城市", placeholder="如: 北京, 杭州")
    filters["education"] = st.selectbox("最高学历", options=["不限", "大专", "本科", "硕士", "博士"], index=0)
    filters["min_experience"] = st.slider("最低工作年限", min_value=0, max_value=20, value=0, step=1)
    if skills:
        filters["skill"] = st.multiselect("技能要求", options=skills, default=[], help="选择必须匹配的技能")
    else:
        filters["skill"] = st.text_input("技能要求", placeholder="如: Python, React")
    col1, col2 = st.columns(2)
    with col1:
        filters["is_985"] = st.checkbox("985 院校", value=False)
    with col2:
        filters["is_211"] = st.checkbox("211 院校", value=False)
    col1, col2 = st.columns(2)
    with col1:
        if st.button("应用筛选", type="primary", use_container_width=True):
            st.session_state.search_filters = filters
    with col2:
        if st.button("重置", use_container_width=True):
            st.session_state.search_filters = {}
            st.rerun()
    return filters


def render_stats_card(label, value, delta=None, help_text=None):
    """Render a styled statistics metric card."""
    st.metric(label=label, value=value, delta=delta, help=help_text)


def export_to_excel(data, filename="export.xlsx"):
    """Export list of dicts to Excel bytes using openpyxl."""
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "导出数据"
        if not data:
            ws.cell(row=1, column=1, value="无数据")
            out = BytesIO()
            wb.save(out)
            out.seek(0)
            return out.getvalue()
        hf = Font(bold=True, color="FFFFFF", size=11)
        hfill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        ha = Alignment(horizontal="center", vertical="center")
        tb = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
        headers = list(data[0].keys())
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = hf; cell.fill = hfill; cell.alignment = ha; cell.border = tb
        for i, row_data in enumerate(data):
            row = i + 2
            for col, key in enumerate(headers, 1):
                val = row_data.get(key, "")
                if isinstance(val, list):
                    val = ", ".join(str(v) for v in val)
                cell = ws.cell(row=row, column=col, value=val)
                cell.border = tb
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15
        out = BytesIO()
        wb.save(out)
        out.seek(0)
        return out.getvalue()
    except Exception as e:
        st.error(f"导出Excel失败: {e}")
        return None